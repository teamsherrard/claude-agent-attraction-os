#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
render_doc.py — Agent Attraction Brain universal document renderer, BOOK edition (v2 type scale).

Turns the house structured text (CAPS section bands, • bullets, "Label:" lead-ins,
sub-bands, pipe tables, cues, the credibility stamp) into a styled .docx — Arial,
near-black, one neutral standard for every member — PLUS a premium long-form "book"
mode for multi-chapter deliverables like the Agent Attraction Brain Book.

BOOK MODE — how it activates (the trigger)
------------------------------------------
Book mode turns on when EITHER:
  • the input contains a [[TOC]] ... [[/TOC]] block   (the primary, content-driven
    trigger — skills opt in simply by writing a contents block; no CLI change), OR
  • the --book CLI flag is passed                      (for cover + chapter breaks on a
    book that has no contents page).
Docs with neither render in LEGACY mode: title block on page 1, no cover break, no
contents page, no bookmarks, no chapter page breaks — the one-off document look.

REAL HEADING STYLES (v2)
------------------------
Every heading is a real Word style so Word's navigation pane and the Google Docs
outline work:
  • PART / CHAPTER band titles, the CONTENTS title and legacy CAPS bands → Heading 1
  • sub-bands ("──── Label ────") and ALL-CAPS sub-labels              → Heading 2
  • the cover / title-block title                                        → Title
The built-in styles are overridden to the house look (Arial, near-black 111111, no
theme colour, no underline, keep-with-next on) so nothing turns blue or serif. The
contents page stays a manual linked list — Word TOC fields do not update in Docs.

TYPE SCALE (v2)
---------------
  body 11pt · line spacing 1.2 · space after 6pt · widow/orphan control on
  cover title 30pt · part title 24pt tracked caps + 10pt accent kicker
  chapter title 20pt bold tracked caps + 9.5pt muted kicker · legacy band 14pt caps
  section label (Heading 2) 12.5pt bold · table text 10pt (header bold on a light
  accent tint, rows white / F7F7F8, hairline E4E4E6 horizontal rules + outer frame,
  header row repeats across pages) · TOC entry 11pt + 9.5pt muted summary
  eyebrow/kicker 9.5pt muted tracked · ">> " callout 11pt upright, indented, 2.25pt
  accent bar on the left over a very light tint

HEADER + FOOTER (both modes; "different first page" — the cover / title page has neither)
  header right: the document title, 8.5pt muted
  footer left: the subtitle (member line) + the eyebrow, 8.5pt muted
  footer right: "Page X of Y" (PAGE + NUMPAGES fields, cached "1" so viewers that
  never update fields still show a number; both become live in Google Docs)

ACCENT COLOUR
-------------
  --accent HEX (default 111111) colours the cover rule, the PART kickers, the callout
  bar and the table-header tint (the accent at ~12% over white). If the accent is too
  light to read on white (contrast < 4.5:1) text uses fall back to 111111; the tints
  keep the accent. Pass the member's primary brand colour when they have one.

BOOK GRAMMAR
------------
1. COVER PAGE — the title block (eyebrow → 30pt title → accent hairline →
   subtitle/byline/meta lines) with top air, then a page break before the body.

2. CONTENTS block:
       [[TOC]]
       PART I — WHO YOU ARE
       CHAPTER 3 — MARKET INTELLIGENCE :: one-line summary
       [[/TOC]]
   Renders a "CONTENTS" page: each entry is an internal hyperlink (w:hyperlink
   w:anchor=...) to the bookmark on the matching body heading, with the summary
   underneath in 9.5pt grey; PART rows render bold. No page numbers (not computable
   at render time; the links are the navigation). Page break after.
   Entries may be written WITH or WITHOUT the "CHAPTER n —"/"PART X —" prefix —
   matching normalizes both sides (strip prefix, case, punctuation).
   FAIL-SOFT: an entry with no matching heading renders as plain text and emits
       WARNING: TOC entry "<title>" has no matching heading - rendered as plain text
   on stderr (the verify gate greps for WARNING).

3. CHAPTER / PART bands — a CAPS band heading matching
       ^(PART [IVXLC]+|CHAPTER \\d+) <dash> TITLE
   renders a small tracked kicker line ("CHAPTER 3") above the Heading 1 title and
   takes a PAGE BREAK BEFORE it — except when it is the first body element (the
   cover/TOC already broke). EVERY band heading is bookmarked with a deterministic
   slug. A PART page also gets "IN THIS PART": a linked list of that part's chapters
   (number · title · the contents summary), built from the [[TOC]] data, emitted
   after the part's intro paragraph(s) — a part page is never near-empty.

4. ">> " CALLOUT — a line starting ">> " renders as the key-insight block (bold lead
   when the text starts with a "Label:" pattern). Book mode only — in legacy docs
   ">>" lines keep their script-cue rendering; script cue heads (ON SCREEN / PAUSE /
   FACT: / "[") stay cues in both modes.

BOOKMARK SLUG RULE (deterministic)
----------------------------------
  norm  = heading text → strip leading "CHAPTER n"/"PART X" + dash → UPPER
          → every non-[A-Z0-9 ] char becomes a space → collapse whitespace
  slug  = "bm_" + norm.lower() with non-alphanumerics collapsed to "_",
          truncated to 32 chars after the prefix; duplicate slugs get _2, _3 …
          in document order.  (≤ 40 chars, starts with a letter — Word-valid.)
  The TOC matches entries to headings by comparing `norm` on both sides.

SUB-BAND CAP
------------
  "──── Label ────" labels are capped at 80 characters. A longer label does NOT
  parse as a sub-band: it falls through to body text and renders WITH ITS LITERAL
  DASHES — a visible failure — and the renderer prints
       WARNING: sub-band label is N characters (cap 80) - rendered as body text with literal dashes: "..."
  on stderr so the verify gate catches it. Shorten the label and re-render.

GOOGLE DOCS CONVERSION (.docx uploaded to Drive, opened in Docs)
----------------------------------------------------------------
SURVIVES: heading styles (→ the Docs outline), bookmarks + internal w:anchor
hyperlinks, page breaks and w:pageBreakBefore, table shading/borders, the repeating
header row, header/footer with PAGE and NUMPAGES fields, paragraph borders/shading
on callouts, bold/size/colour, cell margins.
KNOWN RISKS: letter-tracking (w:spacing) is dropped by Docs (cosmetic — headings
simply render untracked); keep-with-next is advisory in Docs. All degrade
gracefully; nothing breaks navigation or content.

Usage:
  python3 render_doc.py INPUT.txt "OUTPUT.docx"
      [--title "Doc Title"] [--subtitle "Agent · City"] [--eyebrow "KICKER"]
      [--accent 1F3A5F]   # member's primary colour; default 111111 (near-black)
      [--book]            # force book mode without a [[TOC]] block

Requires python-docx.
"""
import sys, re, argparse

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.stderr.write(
        "RENDERER-UNAVAILABLE: python-docx is not installed in this environment.\n"
        "DO NOT attempt to install it. DO NOT run pip, and do not retry this command --\n"
        "in a sandbox a package install can block for many minutes and looks like a hang.\n"
        "Fall back immediately: save the structured text as a .md/.txt file instead, upload\n"
        "THAT, and tell the agent in one plain line that the styled version needs the\n"
        "renderer. Never let this stop the delivery.\n")
    sys.exit(2)

# ---------- the one house style (neutral PREMIUM — accent only where the flag says) ----------
INK_HEX = "111111"
INK    = RGBColor(0x11, 0x11, 0x11)  # title + headings — near-black
BODY   = RGBColor(0x1A, 0x1A, 0x1A)  # body copy — near-black, crisp and legible
MUTED  = RGBColor(0x60, 0x63, 0x67)  # refined mid-grey — eyebrow / subtitle / meta / small print
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
RULE   = "D7DADD"                     # hairline section rule
ACCENT = INK_HEX                      # default accent (cover rule, part kickers, callout bar, header tint)
TBL_LINE = "E4E4E6"                   # table hairlines
ALT_FILL = "F7F7F8"                   # table body alternate row
COOL_FILL = "F7F8F9"                  # legacy example box
FONT   = "Arial"

# type scale (points)
BODY_PT, BODY_LINE, BODY_AFTER = 11, 1.2, 6
TABLE_PT = 10
SMALL_PT = 9.5                        # kicker / eyebrow / TOC summary / cue
CHAPTER_PT, PART_PT, LEGACY_BAND_PT = 20, 24, 14
SUBHEAD_PT = 12.5
COVER_TITLE_PT = 30
HF_PT = 8.5                           # header / footer

# OOXML child order (so inserted elements validate): successors of each element we add by hand
_PBDR_SUCC = ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
              'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
              'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
              'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId',
              'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
_PSHD_SUCC = _PBDR_SUCC[1:]
_RSPACING_SUCC = ('w:w', 'w:kern', 'w:position', 'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect', 'w:bdr',
                  'w:shd', 'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang', 'w:eastAsianLayout',
                  'w:specVanish', 'w:oMath')
_TCBDR_SUCC = ('w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark',
               'w:headers', 'w:cellIns', 'w:cellDel', 'w:cellMerge', 'w:tcPrChange')
_TCSHD_SUCC = _TCBDR_SUCC[1:]
_TCMAR_SUCC = _TCBDR_SUCC[3:]
_TBLBDR_SUCC = ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription')
_CANTSPLIT_SUCC = ('w:trHeight', 'w:tblHeader', 'w:tblCellSpacing', 'w:jc', 'w:hidden', 'w:ins', 'w:del', 'w:trPrChange')
_TBLHEADER_SUCC = _CANTSPLIT_SUCC[2:]

# ---------- colour helpers ----------
def _hex(s, default=ACCENT):
    s = (s or "").strip().lstrip('#').upper()
    if re.fullmatch(r'[0-9A-F]{6}', s): return s
    if re.fullmatch(r'[0-9A-F]{3}', s): return "".join(c * 2 for c in s)
    if s: sys.stderr.write('WARNING: --accent "%s" is not a hex colour - using %s\n' % (s, default))
    return default

def _lum(hexc):
    def ch(i):
        c = int(hexc[i:i + 2], 16) / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * ch(0) + 0.7152 * ch(2) + 0.0722 * ch(4)

def _contrast_on_white(hexc): return 1.05 / (_lum(hexc) + 0.05)

def _tint(hexc, pct):
    """The colour at `pct` strength over white (0.12 → a light tint)."""
    return "".join("%02X" % int(round(255 - (255 - int(hexc[i:i + 2], 16)) * pct)) for i in (0, 2, 4))

# ---------- run / paragraph helpers ----------
def _font(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)

def _track(run, val=0):
    if not val: return run
    rpr = run._element.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(val))
    rpr.insert_element_before(sp, *_RSPACING_SUCC)
    return run

def _run(p, text, size=BODY_PT, color=BODY, bold=False, italic=False, track=0):
    r = p.add_run(text); _font(r); r.font.size = Pt(size)
    r.font.color.rgb = color; r.bold = bold; r.italic = italic
    if track: _track(r, track)
    return r

def _sp(p, before=0, after=BODY_AFTER, line=BODY_LINE):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    return p

def _keep(p, next_=True, together=False, widow=True):
    pf = p.paragraph_format
    if next_: pf.keep_with_next = True
    if together: pf.keep_together = True
    if widow: pf.widow_control = True
    return p

def _border(p, color=RULE, sz=12, space=5, side='bottom'):
    pPr = p._p.get_or_add_pPr()
    b = pPr.find(qn('w:pBdr'))
    if b is None:
        b = OxmlElement('w:pBdr'); pPr.insert_element_before(b, *_PBDR_SUCC)
    bt = OxmlElement('w:' + side)
    bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), str(sz))
    bt.set(qn('w:space'), str(space)); bt.set(qn('w:color'), color)
    b.append(bt)

def _pshade(p, hexc):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexc)
    pPr.insert_element_before(shd, *_PSHD_SUCC)

def _cell_bg(cell, hexc):
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexc)
    cell._tc.get_or_add_tcPr().insert_element_before(shd, *_TCSHD_SUCC)

def _cell_borders(cell, color=TBL_LINE, sz=4, left=True, right=True):
    """Hairline top/bottom; left/right only where asked (the outer frame) — no inner vertical lines."""
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for s, on in (('top', True), ('left', left), ('bottom', True), ('right', right)):
        e = OxmlElement(f'w:{s}')
        if on:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    tcPr.insert_element_before(b, *_TCBDR_SUCC)

def _cell_margins(cell, t=80, b=80, l=130, r=130):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for s, v in (('top', t), ('start', l), ('bottom', b), ('end', r)):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.insert_element_before(m, *_TCMAR_SUCC)

def _tbl_borders(t, color=TBL_LINE, sz=4):
    """Table-level borders: outer frame + inside horizontal rules, no inside vertical rules."""
    b = OxmlElement('w:tblBorders')
    for s, val in (('top', 'single'), ('left', 'single'), ('bottom', 'single'), ('right', 'single'),
                   ('insideH', 'single'), ('insideV', 'nil')):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), val)
        if val != 'nil':
            e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    t._tbl.tblPr.insert_element_before(b, *_TBLBDR_SUCC)

def _row_props(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement('w:cantSplit'); trPr.insert_element_before(cs, *_CANTSPLIT_SUCC)
    if header:
        th = OxmlElement('w:tblHeader'); trPr.insert_element_before(th, *_TBLHEADER_SUCC)

def _hyperlink(p, anchor, text, size=BODY_PT, color=INK_HEX, bold=False):
    """Internal hyperlink run: <w:hyperlink w:anchor="...">. Underlined (the only
    underline in the house style)."""
    h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), FONT)
    rPr.append(rf)
    if bold: rPr.append(OxmlElement('w:b'))
    c = OxmlElement('w:color'); c.set(qn('w:val'), color); rPr.append(c)
    for tag in ('w:sz', 'w:szCs'):
        e = OxmlElement(tag); e.set(qn('w:val'), str(int(round(size * 2)))); rPr.append(e)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t); h.append(r)
    p._p.append(h)

class Doc:
    def __init__(self, accent=ACCENT):
        self.book = False           # set by render() when book mode is active
        self.accent = _hex(accent)
        # text uses (part kicker) need contrast on white; tints keep the accent regardless
        self.accent_text = self.accent if _contrast_on_white(self.accent) >= 4.5 else INK_HEX
        self.head_tint = _tint(self.accent, 0.12)       # table header row
        self.callout_tint = _tint(self.accent, 0.06)    # ">> " callout block
        self.d = Document()
        self._house_styles()
        s = self.d.sections[0]
        s.top_margin = Inches(1.0); s.bottom_margin = Inches(0.9)
        s.left_margin = Inches(1.0); s.right_margin = Inches(1.0)
        self.usable = s.page_width - s.left_margin - s.right_margin
        self._bmid = 0

    # ---- house overrides of the built-in styles ----------------------------
    def _style_font(self, st, size, bold, color=INK):
        f = st.font
        f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = False; f.underline = False
        f.color.rgb = color
        rPr = st.element.get_or_add_rPr()
        rf = rPr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            rf.attrib.pop(qn(a), None)                      # no theme fonts → never serif / Calibri
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), FONT)
        c = rPr.find(qn('w:color'))
        if c is not None:
            for a in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
                c.attrib.pop(qn(a), None)                   # no theme colour → never blue
        szcs = rPr.find(qn('w:szCs'))
        if szcs is not None: szcs.set(qn('w:val'), str(int(round(size * 2))))

    def _house_styles(self):
        st = self.d.styles
        n = st['Normal']; self._style_font(n, BODY_PT, False, BODY)
        pf = n.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(BODY_AFTER)
        pf.line_spacing = BODY_LINE; pf.widow_control = True
        for name, size, before, after, line in (('Heading 1', CHAPTER_PT, 0, 10, 1.02),
                                                ('Heading 2', SUBHEAD_PT, 11, 3, 1.15)):
            h = st[name]; self._style_font(h, size, True, INK)
            pf = h.paragraph_format
            pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
            pf.keep_with_next = True; pf.keep_together = True; pf.widow_control = True
        try:
            t = st['Title']; self._style_font(t, COVER_TITLE_PT, True, INK)
            pPr = t.element.pPr
            if pPr is not None:
                for tag in ('w:pBdr', 'w:contextualSpacing'):
                    e = pPr.find(qn(tag))
                    if e is not None: pPr.remove(e)
            pf = t.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(6); pf.line_spacing = 1.04
            pf.keep_with_next = True; pf.widow_control = True
        except KeyError:
            pass
        for name in ('Heading 1 Char', 'Heading 2 Char', 'Title Char'):
            try:
                self._style_font(st[name], {'Heading 1 Char': CHAPTER_PT, 'Heading 2 Char': SUBHEAD_PT}.get(name, COVER_TITLE_PT), True, INK)
            except KeyError:
                pass

    # ---- bookmarks --------------------------------------------------------
    def _bookmark(self, p, name):
        self._bmid += 1
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(self._bmid)); start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(self._bmid))
        el = p._p
        pPr = el.find(qn('w:pPr'))
        if pPr is not None: pPr.addnext(start)
        else: el.insert(0, start)
        el.append(end)

    # ---- headings ---------------------------------------------------------
    def heading(self, text, slug=None):
        """Plain CAPS band (legacy docs, or a book band without a PART/CHAPTER prefix): Heading 1."""
        p = self.d.add_paragraph(style='Heading 1'); _sp(p, 22, 8, 1.0); _keep(p, together=True)
        _run(p, text.upper(), LEGACY_BAND_PT, INK, bold=True, track=70)
        _border(p, RULE, 6, 7)
        if slug: self._bookmark(p, slug)

    def chapter_heading(self, kicker, title_text, slug=None, page_break=True):
        """CHAPTER/PART band: small tracked kicker (PART kickers in the accent), then the
        Heading 1 title with the hairline rule. Page-breaks before itself unless it is
        the first body element. Kicker + title are keep-with-next: a band never ends a page."""
        is_part = kicker.upper().startswith('PART')
        p = self.d.add_paragraph(); _sp(p, 0, 5, 1.0); _keep(p)
        if page_break: p.paragraph_format.page_break_before = True
        if is_part:
            _run(p, kicker.upper(), 10, RGBColor.from_string(self.accent_text), bold=True, track=140)
        else:
            _run(p, kicker.upper(), SMALL_PT, MUTED, bold=True, track=140)
        hp = self.d.add_paragraph(style='Heading 1'); _sp(hp, 0, 10, 1.02); _keep(hp, together=True)
        _run(hp, title_text.upper(), PART_PT if is_part else CHAPTER_PT, INK, bold=True, track=60 if is_part else 50)
        _border(hp, RULE, 6, 7)
        if slug: self._bookmark(hp, slug)

    def subheading(self, text, note=""):
        p = self.d.add_paragraph(style='Heading 2'); _sp(p, 11, 3, 1.15); _keep(p, together=True)
        _run(p, text.strip(), SUBHEAD_PT, INK, bold=True, track=15)
        if note: _run(p, "   " + note, SMALL_PT, MUTED, italic=True)

    def part_list(self, items):
        """"IN THIS PART" — a linked list of the part's chapters (number · title · summary),
        from the contents data, so a part opener page is never near-empty."""
        k = self.d.add_paragraph(); _sp(k, 16, 4, 1.0); _keep(k)
        _run(k, "IN THIS PART", SMALL_PT, MUTED, bold=True, track=140)
        for label, slug, summary in items:
            p = self.d.add_paragraph(); _sp(p, 2, 4, 1.15); _keep(p, next_=False)
            p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.3)
            if slug: _hyperlink(p, slug, label, size=BODY_PT)
            else:    _run(p, label, BODY_PT, INK)
            if summary: _run(p, " · " + summary, SMALL_PT, MUTED)

    # ---- body elements ----------------------------------------------------
    def body(self, text, size=BODY_PT, keep_next=False):
        p = self.d.add_paragraph(); _sp(p, 0, BODY_AFTER); _keep(p, next_=keep_next)
        m = re.match(r'^([A-Z][A-Za-z0-9 /&→\-\(\)]{2,40}?):\s\s?(.*)$', text)
        if m and len(m.group(1)) < 38:
            _run(p, m.group(1) + ":  ", size, INK, bold=True); _run(p, m.group(2), size)
        else:
            _run(p, text, size)

    def bullet(self, text, lead=None):
        p = self.d.add_paragraph(); _sp(p, 0, 4); _keep(p, next_=False)
        p.paragraph_format.left_indent = Inches(0.28); p.paragraph_format.first_line_indent = Inches(-0.18)
        _run(p, "•  ", BODY_PT, INK, bold=True)
        if lead: _run(p, lead, BODY_PT, INK, bold=True)
        _run(p, text, BODY_PT)

    def numbered(self, n, lead, text):
        p = self.d.add_paragraph(); _sp(p, 4, 5); _keep(p, next_=False)
        p.paragraph_format.left_indent = Inches(0.4); p.paragraph_format.first_line_indent = Inches(-0.4)
        _run(p, f"{n}.  ", 12, INK, bold=True)
        if lead: _run(p, lead + " ", BODY_PT, INK, bold=True)
        _run(p, text, BODY_PT)

    def cue(self, text):
        p = self.d.add_paragraph(); _sp(p, 1, 3)
        p.paragraph_format.left_indent = Inches(0.28)
        _run(p, text, SMALL_PT, MUTED, italic=True)

    def callout(self, text, label=None):
        # legacy example box (script structure branch) — single-cell shaded table
        t = self.d.add_table(rows=1, cols=1); t.autofit = False
        c = t.rows[0].cells[0]; c.width = self.usable; c.text = ""
        p = c.paragraphs[0]; _sp(p, 2, 2, 1.12)
        if label: _run(p, label + "  ", TABLE_PT, INK, bold=True)
        _run(p, text, TABLE_PT, BODY, italic=True)
        _cell_bg(c, COOL_FILL); _cell_borders(c, "E6E3DC"); _cell_margins(c, 90, 90, 150, 150)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def insight(self, text, label=None):
        """">> " key-insight callout: an indented block with a 2.25pt accent bar on the
        left over a very light accent tint; 11pt upright; keep-with-next OFF."""
        p = self.d.add_paragraph(); _sp(p, 6, 8, BODY_LINE); _keep(p, next_=False, together=True)
        pf = p.paragraph_format; pf.left_indent = Inches(0.35); pf.right_indent = Inches(0.15)
        _border(p, self.accent, 18, 8, side='left')          # 18 eighths = 2.25pt
        _pshade(p, self.callout_tint)
        if label: _run(p, label + ":  ", BODY_PT, INK, bold=True)
        _run(p, text, BODY_PT, BODY)

    def table(self, headers, rows, widths):
        t = self.d.add_table(rows=1, cols=len(headers)); t.autofit = True
        tw = t._tbl.tblPr.find(qn('w:tblW'))
        if tw is not None:
            tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(self.usable / 635)))
        _tbl_borders(t)
        last = len(headers) - 1
        for i, h in enumerate(headers):
            self._cell(t.rows[0].cells[i], h, TABLE_PT, INK, bold=True, bg=self.head_tint,
                       left=(i == 0), right=(i == last), keep_next=True)
        _row_props(t.rows[0], header=True)                   # repeats across page breaks
        for ri, row in enumerate(rows):
            r = t.add_row(); cells = r.cells; bg = ALT_FILL if ri % 2 else "FFFFFF"
            for i, val in enumerate(row):
                al = 'center' if (i == 0 and len(val) <= 3) else 'left'
                self._cell(cells[i], val, TABLE_PT, BODY, bg=bg, align=al, left=(i == 0), right=(i == last))
            _row_props(r)
        for i, w in enumerate(widths):
            t.columns[i].width = w
            for row in t.rows: row.cells[i].width = w
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def _cell(self, cell, text, size, color, bold=False, align='left', bg=None, left=True, right=True, keep_next=False):
        cell.text = ""; p = cell.paragraphs[0]
        p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
        _sp(p, 1, 1, 1.05); _keep(p, next_=keep_next, widow=False); _run(p, text, size, color, bold)
        if bg: _cell_bg(cell, bg)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _cell_borders(cell, left=left, right=right); _cell_margins(cell)

    # ---- title block (legacy) / cover (book) ------------------------------
    def title_block(self, title, subtitle, meta_lines, eyebrow=None):
        if eyebrow:
            p = self.d.add_paragraph(); _sp(p, 22, 6); _keep(p)
            _run(p, eyebrow.upper(), SMALL_PT, MUTED, bold=True, track=140)
        else:
            self.d.add_paragraph().paragraph_format.space_after = Pt(12)
        p = self.d.add_paragraph(style='Title'); _sp(p, 0, 2, 1.02); _keep(p)
        _run(p, title, COVER_TITLE_PT, INK, bold=True, track=-8)
        if subtitle:
            p = self.d.add_paragraph(); _sp(p, 1, 4); _keep(p)
            _run(p, subtitle, 12.5, MUTED, track=20)
        p = self.d.add_paragraph(); _sp(p, 4, 7); _border(p, self.accent, 8, 7)
        for ml in meta_lines:
            p = self.d.add_paragraph(); _sp(p, 3, 0)
            if ml.lower().startswith("powered by"):
                _run(p, ml.upper(), HF_PT, MUTED, bold=True, track=60)
            else:
                _run(p, ml, SMALL_PT, MUTED)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def cover(self, title, subtitle, meta_lines, eyebrow=None):
        """True cover page: top air → eyebrow → 30pt title → accent hairline →
        byline/date lines → PAGE BREAK before the body."""
        air = self.d.add_paragraph()
        air.paragraph_format.space_before = Pt(150); air.paragraph_format.space_after = Pt(0)
        if eyebrow:
            p = self.d.add_paragraph(); _sp(p, 0, 10); _keep(p)
            _run(p, eyebrow.upper(), SMALL_PT, MUTED, bold=True, track=160)
        p = self.d.add_paragraph(style='Title'); _sp(p, 0, 6, 1.04); _keep(p)
        _run(p, title, COVER_TITLE_PT, INK, bold=True, track=-8)
        rp = self.d.add_paragraph(); _sp(rp, 2, 10); _border(rp, self.accent, 8, 7)
        if subtitle:
            p = self.d.add_paragraph(); _sp(p, 2, 5)
            _run(p, subtitle, 13, MUTED, track=20)
        for ml in meta_lines:
            p = self.d.add_paragraph(); _sp(p, 3, 0)
            if ml.lower().startswith("powered by"):
                _run(p, ml.upper(), HF_PT, MUTED, bold=True, track=60)
            else:
                _run(p, ml, 10, MUTED)
        self.d.add_page_break()

    # ---- contents page ----------------------------------------------------
    def toc_page(self, entries, norm_map):
        p = self.d.add_paragraph(style='Heading 1'); _sp(p, 6, 12, 1.0); _keep(p, together=True)
        _run(p, "CONTENTS", CHAPTER_PT, INK, bold=True, track=90)
        _border(p, RULE, 6, 7)
        for etitle, summary in entries:
            slug = norm_map.get(_norm_title(etitle))
            mp = _CHAP_PREFIX.match(etitle.strip().upper())
            is_part = bool(mp and mp.group(1).startswith('PART'))
            ep = self.d.add_paragraph(); _sp(ep, 10 if is_part else 5, 1 if summary else 4, 1.1)
            _keep(ep, next_=bool(summary) or is_part)
            if slug:
                _hyperlink(ep, slug, etitle, size=BODY_PT, bold=is_part)
            else:
                sys.stderr.write('WARNING: TOC entry "%s" has no matching heading - rendered as plain text\n' % etitle)
                _run(ep, etitle, BODY_PT, INK, bold=is_part)
            if summary:
                sp2 = self.d.add_paragraph(); _sp(sp2, 0, 5, 1.12)
                sp2.paragraph_format.left_indent = Inches(0.18)
                _run(sp2, summary, SMALL_PT, MUTED)
        self.d.add_page_break()

    # ---- header + footer (both modes; the cover / title page has neither) ----
    def _field(self, p, instr, cached="1"):
        def styled():
            r = p.add_run(); _font(r); r.font.size = Pt(HF_PT); r.font.color.rgb = MUTED
            return r
        r = styled(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
        r = styled(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' %s ' % instr; r._r.append(it)
        r = styled(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'separate'); r._r.append(fc)
        r = styled(); t = OxmlElement('w:t'); t.text = cached; r._r.append(t)
        r = styled(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'end'); r._r.append(fc)

    def header_footer(self, title, subtitle=None, eyebrow=None):
        """Different first page (blank header + footer on the cover / title page).
        Header right: the title. Footer: subtitle + eyebrow on the left, "Page X of Y" on
        the right (PAGE + NUMPAGES fields with cached results)."""
        s = self.d.sections[0]
        s.different_first_page_header_footer = True
        s.first_page_header.is_linked_to_previous = False      # explicit, empty
        s.first_page_footer.is_linked_to_previous = False
        h = s.header; h.is_linked_to_previous = False
        hp = h.paragraphs[0] if h.paragraphs else h.add_paragraph()
        hp.text = ""; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _sp(hp, 0, 0, 1.0)
        _run(hp, title or "", HF_PT, MUTED)
        f = s.footer; f.is_linked_to_previous = False
        fp = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        fp.text = ""; fp.alignment = WD_ALIGN_PARAGRAPH.LEFT; _sp(fp, 6, 0, 1.0)
        fp.paragraph_format.tab_stops.add_tab_stop(self.usable, WD_TAB_ALIGNMENT.RIGHT)
        left = "  ·  ".join(x for x in (subtitle, (eyebrow or "").upper() or None) if x)
        _run(fp, left, HF_PT, MUTED)
        _run(fp, "\t", HF_PT, MUTED)
        _run(fp, "Page ", HF_PT, MUTED); self._field(fp, 'PAGE')
        _run(fp, " of ", HF_PT, MUTED); self._field(fp, 'NUMPAGES')

    def stamp(self, text):
        p = self.d.add_paragraph(); _sp(p, 8, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, text.upper(), 9, MUTED, bold=True)

    def footer_note(self, text):
        p = self.d.add_paragraph(); _sp(p, 8, 0); _border(p, "D9D9D9", 6, 6)
        cp = self.d.add_paragraph(); _sp(cp, 4, 0); _run(cp, text, HF_PT, MUTED, italic=True)

    def save(self, path): self.d.save(path)

# ---------- parsing the house-style structured text ----------
_DIV = set("═─—–-")
def is_band(l):
    s = l.strip(); return len(s) > 8 and bool(s) and set(s) <= _DIV
def heavy(l): return bool(l.strip()) and set(l.strip()) == {"═"} and len(l.strip()) > 8
def heading_like(s):
    base = s.split("(")[0]
    letters = [c for c in base if c.isalpha()]
    if not letters: return False
    return sum(c.isupper() for c in letters) / len(letters) >= 0.6 and len(s) <= 80
NUM = re.compile(r'^\s*(\d+)\s')
SUBBAND_CAP = 80
SUBBAND = re.compile(r'^\s*[─═—-]{2,}\s*([^─═]{2,%d}?)\s*[─═—-]{2,}\s*$' % SUBBAND_CAP)
SUBBAND_ANY = re.compile(r'^\s*[─═—-]{2,}\s*([^─═]{2,}?)\s*[─═—-]{2,}\s*$')   # over-cap detector (warning only)

# a dotted-leader metric row: the house style ('Label ........ value') OR an indented row with a short
# 3-dot leader ('   Leads / calls booked ... not tracked yet') — the skills sometimes write the short form
DOTROW = re.compile(r'^\s{2,}\S.*?\s\.{3,}\s+\S')
def _isdot(l):
    return ('....' in l and re.match(r'^\s*\S.*\.{3,}\s*\S', l) is not None) or DOTROW.match(l) is not None
CUE = re.compile(r'^\s*(>>|\[|FACT:|ON SCREEN|PAUSE)')
CHAP = re.compile(r'^(PART\s+[IVXLCDM]+|CHAPTER\s+\d+)\s*[—–:\-]\s*(.+)$')
_CHAP_PREFIX = re.compile(r'^\s*(PART\s+[IVXLCDM]+|CHAPTER\s+\d+)\s*(?:[—–:\-]\s*)?')
_LABEL = re.compile(r'^([A-Z][A-Za-z0-9 /&→\-\(\)]{2,40}?):\s+(.*)$')
_WEEKV = r'^\s*Week\s+(\d+)\s*[·:]\s*(Video|Post)\s*(\d+)\s*—\s*(.*?)\s*\(([^)]*)\)\s*$'
_WEEKP = r'^\s*Week\s+(\d+)\s+—\s+(.*)\((Pillar[^)]*)\)\s*$'

def _starts_table(l):
    """Does this line open one of the table shapes the body loop renders as a table?"""
    s = l.strip()
    if s.startswith("|") and s.endswith("|") and s.count("|") >= 3: return True
    if _isdot(l): return True
    if re.match(r'^\s*#\s+EXACT TITLE', l): return True
    return bool(re.match(_WEEKV, l) or re.match(_WEEKP, l))

def _norm_title(s):
    """Normalize a heading or TOC-entry title for matching: strip CHAPTER/PART
    prefix, uppercase, strip punctuation, collapse whitespace."""
    up = s.strip().upper()
    stripped = _CHAP_PREFIX.sub('', up)
    if not stripped.strip(): stripped = up          # "PART I" alone — keep the words
    stripped = re.sub(r'[^A-Z0-9 ]+', ' ', stripped)
    return re.sub(r'\s+', ' ', stripped).strip()

def _assign_slugs(heading_texts):
    """Deterministic slugs in document order; norm_map maps normalized title →
    slug of the FIRST heading bearing it."""
    slugs, norm_map, used = [], {}, {}
    for h in heading_texts:
        norm = _norm_title(h)
        base = re.sub(r'[^a-z0-9]+', '_', norm.lower()).strip('_')[:32] or 'section'
        name = 'bm_' + base
        k = used.get(name, 0) + 1; used[name] = k
        if k > 1: name = f'{name[:36]}_{k}'
        slugs.append(name)
        norm_map.setdefault(norm, name)
    return slugs, norm_map

def _extract_toc(lines):
    """Pull the [[TOC]] block out of the line stream. Returns (lines-without-toc,
    entries [(title, summary)], has_toc)."""
    out, entries, has, in_toc = [], [], False, False
    for l in lines:
        s = l.strip()
        if s == '[[TOC]]': has = True; in_toc = True; continue
        if s == '[[/TOC]]': in_toc = False; continue
        if in_toc:
            if not s: continue
            title, sep, summ = s.partition('::')
            entries.append((title.strip(), summ.strip() if sep else ''))
            continue
        out.append(l)
    return out, entries, has

def _part_lists(entries, norm_map):
    """Group the contents entries under their PART rows → {part norm: [(label, slug, summary)]}
    for the "IN THIS PART" list on each part opener page."""
    parts, cur, seq = {}, None, 0          # seq runs across parts, so prefix-less rows number like CHAPTER n
    for title, summary in entries:
        up = title.strip().upper()
        mp = _CHAP_PREFIX.match(up)
        if mp and mp.group(1).startswith('PART'):
            cur = _norm_title(title); parts[cur] = []; continue
        if cur is None: continue
        seq += 1
        mn = re.match(r'^\s*CHAPTER\s+(\d+)', up)
        num = mn.group(1) if mn else str(seq)
        bare = _CHAP_PREFIX.sub('', up).strip() or up
        parts[cur].append((f"{num} · {bare}", norm_map.get(_norm_title(title)), summary))
    return parts

def _render_body(lines, i, doc, book_mode, collect=None, slug_iter=None, part_lists=None):
    """The section/body loop. collect=list → dry pass recording heading texts;
    slug_iter → real pass consuming precomputed slugs (same order);
    part_lists → the "IN THIS PART" data (real pass, book mode)."""
    n = len(lines)
    last_head = ''
    emitted = False        # has any body element been emitted yet?
    pending_part = None    # part band awaiting its "IN THIS PART" list

    def heading_slug(h):
        if collect is not None:
            collect.append(h); return None
        if slug_iter is not None:
            return next(slug_iter, None)
        return None

    def flush_part():
        nonlocal pending_part
        if pending_part is not None and part_lists:
            items = part_lists.get(pending_part)
            if items: doc.part_list(items)
        pending_part = None

    def section_heading_at(k):
        if is_band(lines[k]) and k+1 < n:
            cand = lines[k+1].strip()
            if cand and not is_band(lines[k+1]):
                wrapped = (k+2 < n and is_band(lines[k+2]))
                if wrapped or heading_like(cand):
                    return cand, (k+3 if wrapped else k+2)
        return None, None

    def table_follows(k):
        """Is the next non-blank line a table opener? (keep a Label: lead with its table)"""
        while k < n and not lines[k].strip(): k += 1
        return k < n and _starts_table(lines[k])

    while i < n:
        if is_band(lines[i]):
            h, nxt = section_heading_at(i)
            if h is not None:
                flush_part()
                slug = heading_slug(h) if book_mode else None
                mch = CHAP.match(h)
                if book_mode and mch:
                    doc.chapter_heading(mch.group(1), mch.group(2).strip(), slug, page_break=emitted)
                    if mch.group(1).upper().startswith('PART') and collect is None:
                        pending_part = _norm_title(h)
                else:
                    doc.heading(h, slug=slug)
                emitted = True; last_head = h; i = nxt; continue
            i += 1; continue            # stray rule — skip
        raw = lines[i]; line = raw.strip()
        if line == "": i += 1; continue
        emitted = True

        # game-plan style tables (safe: only fire on these exact shapes) ----
        if re.match(r'^\s*#\s+EXACT TITLE', raw):
            tcol = raw.find("EXACT TITLE"); icol = raw.find("SEARCH INTENT"); rows = []; i += 1
            while i < n:
                rl = lines[i]
                if rl.strip() == "": i += 1; break
                if not NUM.match(rl): break
                num = rl[:tcol].strip(); rest = rl[tcol:].rstrip()
                gaps = list(re.finditer(r'\s{4,}', rest)) or list(re.finditer(r'\s{2,}', rest))
                if gaps:
                    g = gaps[-1]; ttl, intent = rest[:g.start()].strip(), rest[g.end():].strip()
                    ttl = re.sub(r'\s{2,}', ' ', ttl); intent = re.sub(r'\s{2,}', ' ', intent)
                else:
                    s = max(0, min(icol - tcol, len(rest) - 1))
                    while s > 0 and rest[s] != ' ': s -= 1
                    ttl, intent = rest[:s].strip(), rest[s:].strip()
                rows.append([num, ttl, intent]); i += 1
            doc.table(["#", "Exact video title", "Search intent & lead type"], rows,
                      [Inches(0.4), Inches(4.05), Inches(2.25)])
            continue
        mweekv = re.match(_WEEKV, raw)
        if mweekv:
            rows = []; kind = mweekv.group(2)
            while i < n:
                m = re.match(_WEEKV, lines[i])
                if not m: break
                rows.append([f"Week {m.group(1)}", f"#{m.group(3)}", m.group(4).strip(), m.group(5).strip()]); i += 1
            doc.table(["Week", "#", f"{kind} to publish", "Type"], rows,
                      [Inches(0.7), Inches(0.4), Inches(3.85), Inches(1.75)])
            continue
        mweek = re.match(_WEEKP, raw)
        if mweek:
            rows = []
            while i < n:
                m = re.match(_WEEKP, lines[i])
                if not m: break
                rows.append(["Week " + m.group(1), m.group(2).strip(), m.group(3)]); i += 1
            doc.table(["Week", "Video to publish", "Pillar"], rows,
                      [Inches(0.75), Inches(4.4), Inches(1.55)])
            continue
        if _isdot(raw):
            rows = []
            while i < n and _isdot(lines[i]):
                m = re.match(r'^\s*(.+?)\s*\.{3,}\s*(.+?)\s*$', lines[i])
                if not m: break
                rest = m.group(2).strip(); note = ""
                mm = re.match(r'^(.+?)\s{2,}\((.+)\)\s*$', rest)
                if mm: rest, note = mm.group(1).strip(), mm.group(2).strip()
                rows.append([m.group(1).strip(), rest, note]); i += 1
            lh = (last_head or '').upper()
            if any(k in lh for k in ("COMPETITOR", "OTHER AGENTS")):
                hdr = ["Channel", "Numbers", "Notes"]
            elif any(k in lh for k in ("SCORECARD", "AT A GLANCE", "YOUR NUMBERS")):
                hdr = ["Metric", "This window", "vs last / note"]
            elif any(k in lh for k in ("GOAL", "MATH", "TARGET", "KPI")):
                hdr = ["Metric", "Target", "Why it matters"]
            else:
                hdr = ["Item", "Result", "Note"]
            doc.table(hdr, rows, [Inches(2.1), Inches(2.25), Inches(2.35)])
            continue

        # generic pipe table:  | Head 1 | Head 2 |  /  | --- | --- |  /  | a | b |
        if line.startswith("|") and line.endswith("|") and line.count("|") >= 3:
            rows = []
            while i < n:
                l = lines[i].strip()
                if not (l.startswith("|") and l.endswith("|")): break
                cells = [c.strip() for c in l.strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                rows = [r + [""] * (ncol - len(r)) for r in rows]
                # column widths follow the longest cell in each column (min 4 chars, cap 60), so a title
                # column gets the room and a '#' or 'Views' column stays narrow; floor at 0.55in
                lens = [max((len(str(r[c])) for r in rows), default=4) for c in range(ncol)]
                wts = [max(4, min(l, 60)) for l in lens]; tot = float(sum(wts))
                widths = [max(int(Inches(0.55)), int(doc.usable * w / tot)) for w in wts]
                scale = doc.usable / float(sum(widths)); widths = [int(w * scale) for w in widths]
                doc.table(rows[0], rows[1:] if len(rows) > 1 else [], widths)
            continue

        # sub-band  "──── LABEL ────"  (label ≤ 80 chars; longer falls through WITH its dashes + a warning)
        msb = SUBBAND.match(raw)
        if msb: doc.subheading(msb.group(1)); last_head = msb.group(1); i += 1; continue
        mover = SUBBAND_ANY.match(raw)
        if mover and collect is None:          # warn once (the real pass), not again in the dry pass
            sys.stderr.write('WARNING: sub-band label is %d characters (cap %d) - rendered as body text with literal dashes: "%s"\n'
                             % (len(mover.group(1)), SUBBAND_CAP, mover.group(1)[:60]))
        # bullets
        if line.startswith("•"):
            bt = line[1:].strip()
            mm = re.match(r'^([^:]{3,60}):\s+(.*)$', bt)
            if mm and ("—" in mm.group(1) or len(mm.group(1)) < 45):
                doc.bullet(mm.group(2), lead=mm.group(1) + ":  ")
            else:
                doc.bullet(bt)
            i += 1; continue
        # numbered list item "1.  TEXT"  (gather indented continuations)
        mnum = re.match(r'^(\d+)\.\s+(.*)$', line)
        if mnum and 1 <= int(mnum.group(1)) <= 99:
            txt = mnum.group(2); i += 1
            while (i < n and lines[i].strip() and lines[i].startswith("   ")
                   and not re.match(r'^\s*\d+\.\s', lines[i]) and not is_band(lines[i])
                   and not SUBBAND.match(lines[i]) and not CUE.match(lines[i].strip())):
                txt += " " + lines[i].strip(); i += 1
            ml = re.match(r'^([A-Z][A-Z0-9 ,\-–—/&]{3,55}?[\.\-—])\s*(.*)$', txt)
            if ml: doc.numbered(mnum.group(1), ml.group(1), ml.group(2))
            else:  doc.numbered(mnum.group(1), None, txt)
            continue
        # 4-part-structure style label "HOOK (0:00) — text" + example
        mstruct = re.match(r'^([A-Z][A-Z ]{2,22}?)(\s*\([^)]*\))?\s*—\s*(.+)$', line)
        if mstruct and mstruct.group(1).strip() in (
                "HOOK", "PRIMARY CTA", "SECONDARY CTA", "BODY", "CONTENT", "EARLY CTA", "CLOSING CTA",
                "EARLY WARM CTA", "SOFT MID CTA", "MID CTA", "REASSURANCE", "FINAL CTA", "INTRO"):
            label = mstruct.group(1) + (" " + mstruct.group(2) if mstruct.group(2) else "")
            txt = mstruct.group(3); i += 1
            eg = ""
            _SL = re.compile(r'^\s*([A-Z][A-Z ]{2,22}?)(\s*\([^)]*\))?\s*—\s')
            while (i < n and lines[i].strip() and lines[i].startswith("   ")
                   and not is_band(lines[i]) and not _SL.match(lines[i])):
                eg += (" " if eg else "") + lines[i].strip(); i += 1
            p = doc.d.add_paragraph(); _sp(p, 6, 2); _keep(p, next_=bool(eg))
            _run(p, label, BODY_PT, INK, bold=True); _run(p, "  —  " + txt, BODY_PT, BODY)
            if eg: doc.callout(eg.replace("e.g. ", ""), label="Example" if eg.startswith("e.g.") else None)
            continue
        # ">> " key-insight callout (BOOK grammar; legacy docs keep the cue look).
        # Script cue heads stay cues even in book mode — a book embedding script
        # fragments must not turn ">> ON SCREEN — ..." into an insight box.
        if book_mode and line.startswith(">> ") and not re.match(r'^(ON SCREEN|PAUSE|FACT:|\[)', line[3:].strip()):
            txt = line[3:].strip()
            mlab = _LABEL.match(txt)
            if mlab and len(mlab.group(1)) < 38:
                doc.insight(mlab.group(2), label=mlab.group(1))
            else:
                doc.insight(txt)
            i += 1; continue
        # cue lines (>> ON SCREEN, [PAUSE], FACT:)
        if CUE.match(line): doc.cue(line); i += 1; continue
        # ALL-CAPS sub-label (optionally with a lowercase parenthetical)
        mlbl = re.match(r'^([A-Z][A-Z0-9 &/\-]{5,42}?)(\s*\(.*\))?:?\s*$', line)
        if mlbl and mlbl.group(1).strip().count(' ') >= 1 and not NUM.match(raw):
            doc.subheading(mlbl.group(1).strip(), mlbl.group(2).strip() if mlbl.group(2) else "")
            last_head = mlbl.group(1).strip()
            i += 1; continue
        # the credibility stamp
        if line.lower().startswith("powered by mike sherrard"):
            doc.stamp(line); i += 1; continue
        if line.lower().startswith("compliance"):
            doc.footer_note(line); i += 1; continue
        # default paragraph — a "Label:" lead (or any short lead line) stays with the table that follows it
        is_label = bool(_LABEL.match(line))
        doc.body(line, keep_next=(is_label or len(line) <= 90) and table_follows(i + 1)); i += 1
    flush_part()

def render(lines, doc, title=None, subtitle=None, eyebrow=None, book=False):
    lines, toc_entries, has_toc = _extract_toc(lines)
    book_mode = bool(book or has_toc)
    doc.book = book_mode
    n = len(lines); i = 0
    # ----- title block: everything before the first band -----
    head = []
    while i < n and not is_band(lines[i]) and not SUBBAND.match(lines[i]) and len(head) < 8:
        if lines[i].strip(): head.append(lines[i].strip())
        i += 1
    t = title or (head[0] if head else "Document")
    meta = head[1:]
    sub = subtitle
    if sub is None and meta and "·" in meta[0] and not meta[0].lower().startswith(("prepared", "powered")):
        sub, meta = meta[0], meta[1:]
    elif sub and meta and meta[0].strip() == sub.strip():
        meta = meta[1:]            # --subtitle repeats the input's byline → render it once

    if not book_mode:
        # legacy path — one-off documents: title block, body, header/footer from page 2
        doc.title_block(t, sub, meta, eyebrow=eyebrow)
        _render_body(lines, i, doc, False)
        doc.header_footer(t, sub, eyebrow)
        return

    # ----- book path: cover → (dry pass for slugs) → contents → body (+ part lists) → header/footer -----
    doc.cover(t, sub, meta, eyebrow=eyebrow)
    collected = []
    _render_body(lines, i, Doc(doc.accent), True, collect=collected)   # dry: discover headings in emission order
    slugs, norm_map = _assign_slugs(collected)
    part_lists = _part_lists(toc_entries, norm_map) if has_toc else None
    if has_toc:
        doc.toc_page(toc_entries, norm_map)
    _render_body(lines, i, doc, True, slug_iter=iter(slugs), part_lists=part_lists)
    doc.header_footer(t, sub, eyebrow)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input"); ap.add_argument("output")
    ap.add_argument("--title", default=None); ap.add_argument("--subtitle", default=None)
    ap.add_argument("--eyebrow", default=None)
    ap.add_argument("--accent", default=ACCENT,
                    help="Accent hex colour (e.g. 1F3A5F) for the cover rule, PART kickers, the callout bar "
                         "and the table-header tint. Default 111111 (near-black). A colour too light to read "
                         "on white falls back to 111111 for text; the tints keep it.")
    ap.add_argument("--book", action="store_true",
                    help="Force book mode (cover page break, chapter page breaks, footer page numbers) "
                         "even without a [[TOC]] block. A [[TOC]] block turns book mode on by itself.")
    a = ap.parse_args()
    lines = open(a.input, encoding="utf-8").read().split("\n")
    doc = Doc(accent=a.accent); render(lines, doc, a.title, a.subtitle, a.eyebrow, book=a.book); doc.save(a.output)
    print("rendered:", a.output)

if __name__ == "__main__":
    main()
